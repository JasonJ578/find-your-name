"""把 docs/ 下的 Markdown 教程导出成 Word（.docx）。

用法：
    <内置 python> tools/build_docx.py

输出：
    新手教程/01-新手配置教程.docx
    ...

说明：
    Word 是发给同事阅读的版本；内容以仓库里的 Markdown 为准。
    改了 Markdown 之后重新跑一次本脚本即可。
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "新手教程"

FONT_BODY = "微软雅黑"
FONT_CODE = "Consolas"
FONT_SYMBOL = "Segoe UI Symbol"

TEXT = RGBColor(0x1F, 0x23, 0x29)
MUTED = RGBColor(0x4E, 0x59, 0x69)
LIGHT = RGBColor(0x8A, 0x91, 0x9C)
BRAND = RGBColor(0xFF, 0x82, 0x00)
LINK = RGBColor(0x1A, 0x6F, 0xD4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

DOCS = [
    ("docs/新手配置教程.md", "01-新手配置教程.docx", "新手配置教程", "给零基础同事的配置与工作流指南"),
    ("docs/单页制作流程.md", "02-单页制作流程.docx", "单页制作流程", "从领页到上线的十步"),
    ("docs/页面清单.md", "03-页面清单.docx", "页面清单与归属", "谁负责哪一页"),
    ("docs/项目约定.md", "04-项目约定.docx", "项目约定", "每次让 AI 写页面前先贴这一段"),
    ("docs/变更记录.md", "05-变更记录.docx", "变更记录", "多人共享的项目变更日志"),
    ("docs/开发方案与操作手册.md", "06-开发方案与操作手册.docx", "开发方案与操作手册", "技术形态、AI 工作流、Git 协作与排期"),
    ("docs/Git同步策略.md", "07-Git同步策略.docx", "Git 同步策略", "什么会同步、什么不会、四个人怎么协作"),
]

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
LIST_RE = re.compile(r"^(\s*)([-*]|\d+\.)\s+(.*)$")
SEP_CELL_RE = re.compile(r"^:?-{2,}:?$")


# ---------- 底层格式 ----------

def set_run_font(run, name: str, size: float | None = None,
                 bold: bool | None = None, color: RGBColor | None = None,
                 italic: bool | None = None):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic

    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


def set_style_font(style, name: str, size: float, bold: bool,
                   color: RGBColor | None = None):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color

    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def shade_run(run, fill: str):
    rPr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    rPr.append(shd)


def shade_paragraph(paragraph, fill: str):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def shade_cell(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_bottom_border(paragraph, color: str, size: int = 8):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_left_border(paragraph, color: str, size: int = 18):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    pBdr.append(left)
    pPr.append(pBdr)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


# ---------- 行内格式 ----------

def add_inline(paragraph, text: str, size: float = 10.5,
               color: RGBColor | None = None, italic: bool = False,
               bold: bool = False):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, FONT_BODY, size, bold=True, color=color, italic=italic)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, FONT_CODE, size - 1, color=color, italic=italic)
            shade_run(run, "F2F3F5")
        elif part.startswith("[") and "](" in part:
            match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if match:
                label, url = match.group(1), match.group(2)
                run = paragraph.add_run(label)
                set_run_font(run, FONT_BODY, size, color=LINK, italic=italic)
                run.underline = True
                if url and url != label:
                    tail = paragraph.add_run(f"（{url}）")
                    set_run_font(tail, FONT_BODY, size - 1.5, color=LIGHT, italic=italic)
            else:
                run = paragraph.add_run(part)
                set_run_font(run, FONT_BODY, size, color=color, italic=italic, bold=bold)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, FONT_BODY, size, color=color, italic=italic, bold=bold)


# ---------- 块级元素 ----------

def add_cover(doc, title: str, subtitle: str):
    band = doc.add_paragraph()
    band.paragraph_format.space_before = Pt(0)
    band.paragraph_format.space_after = Pt(60)
    shade_paragraph(band, "FF8200")
    run = band.add_run("  FIND YOUR NAME")
    set_run_font(run, "Arial", 9, bold=True, color=WHITE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(title)
    set_run_font(run, FONT_BODY, 28, bold=True, color=TEXT)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(subtitle)
    set_run_font(run, FONT_BODY, 12, color=MUTED)

    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(24)
    add_bottom_border(line, "FF8200", 12)

    meta = [
        ("文档版本", "v1.0"),
        ("导出日期", "2026-10-24"),
        ("编写", "JasonJ578"),
        ("适用", "Find Your Name 项目组"),
    ]
    for key, value in meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f"{key}：")
        set_run_font(run, FONT_BODY, 10, color=LIGHT)
        run = p.add_run(value)
        set_run_font(run, FONT_BODY, 10, color=MUTED)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(120)
    run = note.add_run("本文件由仓库中的 Markdown 导出；内容以 Markdown 源文件为准。")
    set_run_font(run, FONT_BODY, 8.5, color=LIGHT, italic=True)

    doc.add_page_break()


def add_table(doc, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    size = 8.5 if cols >= 5 else 9.5

    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for col in range(cols):
            text = row[col] if col < len(row) else ""
            cell = cells[col]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.15
            add_inline(paragraph, text, size=size)
            if index == 0:
                for run in paragraph.runs:
                    run.bold = True
                shade_cell(cell, "F2F3F5")

    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "E5E6EB")
        borders.append(el)
    tblPr.append(borders)

    header_tr = table.rows[0]._tr
    trPr = header_tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def add_code(doc, lines: list[str]):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.1
    paragraph.paragraph_format.left_indent = Cm(0.2)
    paragraph.paragraph_format.right_indent = Cm(0.2)
    shade_paragraph(paragraph, "F6F7F9")

    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, FONT_CODE, 9, color=TEXT)


def add_blockquote(doc, lines: list[str]):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.6)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    add_left_border(paragraph, "FF8200", 18)
    text = " ".join(line.lstrip("> ").strip() for line in lines)
    add_inline(paragraph, text, size=10, color=MUTED, italic=True)


def add_separator(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    add_bottom_border(paragraph, "E5E6EB", 6)


def is_cjk(char: str) -> bool:
    return "\u4e00" <= char <= "\u9fff"


def join_paragraph_lines(lines: list[str]) -> str:
    result = lines[0]
    for line in lines[1:]:
        if result and line and is_cjk(result[-1]) and is_cjk(line[0]):
            result += line
        else:
            result += " " + line
    return result


# ---------- 文档构建 ----------

def setup_document() -> Document:
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    set_style_font(normal, FONT_BODY, 10.5, bold=False, color=TEXT)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color, before, after in [
        ("Heading 1", 18, TEXT, 18, 8),
        ("Heading 2", 14, TEXT, 14, 6),
        ("Heading 3", 12, MUTED, 10, 4),
        ("Heading 4", 11, MUTED, 8, 3),
    ]:
        style = doc.styles[name]
        set_style_font(style, FONT_BODY, size, bold=True, color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Bullet 3", "List Number"):
        try:
            style = doc.styles[name]
            set_style_font(style, FONT_BODY, 10.5, bold=False, color=TEXT)
            style.paragraph_format.space_after = Pt(2)
            style.paragraph_format.line_spacing = 1.4
        except KeyError:
            pass

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("Find Your Name · 项目文档")
    set_run_font(run, FONT_BODY, 8, color=LIGHT)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("第 ")
    set_run_font(run, FONT_BODY, 8.5, color=LIGHT)
    add_page_field(fp)
    run = fp.add_run(" 页")
    set_run_font(run, FONT_BODY, 8.5, color=LIGHT)

    return doc


def build(markdown: str, title: str, subtitle: str) -> Document:
    doc = setup_document()
    doc.core_properties.title = title
    doc.core_properties.author = "JasonJ578"
    doc.core_properties.subject = "Find Your Name 项目文档"
    doc.core_properties.comments = "由 tools/build_docx.py 从 Markdown 生成"

    add_cover(doc, title, subtitle)

    lines = markdown.splitlines()
    if lines and HEADING_RE.match(lines[0]) and lines[0].startswith("# "):
        lines = lines[1:]

    paragraph_buffer: list[str] = []

    def flush_paragraph():
        if not paragraph_buffer:
            return
        text = join_paragraph_lines(paragraph_buffer)
        paragraph = doc.add_paragraph()
        add_inline(paragraph, text)
        paragraph_buffer.clear()

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if stripped.startswith("```"):
            flush_paragraph()
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            index += 1
            add_code(doc, code_lines)
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            raw_rows: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                raw_rows.append(lines[index].strip())
                index += 1
            rows: list[list[str]] = []
            for raw in raw_rows:
                cells = [cell.strip() for cell in raw.strip("|").split("|")]
                if cells and all(SEP_CELL_RE.match(cell) for cell in cells):
                    continue
                rows.append(cells)
            add_table(doc, rows)
            continue

        heading = HEADING_RE.match(line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = heading.group(2).strip()
            paragraph = doc.add_heading(level=level)
            run = paragraph.add_run(text)
            set_run_font(run, FONT_BODY, {1: 18, 2: 14, 3: 12, 4: 11}[level],
                         bold=True, color=TEXT if level <= 2 else MUTED)
            if level == 1:
                add_bottom_border(paragraph, "FF8200", 8)
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip())
                index += 1
            add_blockquote(doc, quote_lines)
            continue

        if stripped in ("---", "***", "___"):
            flush_paragraph()
            add_separator(doc)
            index += 1
            continue

        list_match = LIST_RE.match(line)
        if list_match:
            flush_paragraph()
            indent = len(list_match.group(1))
            marker = list_match.group(2)
            text = list_match.group(3).strip()

            if marker in ("-", "*") and text.startswith("[ ] "):
                paragraph = doc.add_paragraph(style="List Bullet" if indent < 2 else "List Bullet 2")
                run = paragraph.add_run("□ ")
                set_run_font(run, FONT_SYMBOL, 10.5, color=MUTED)
                add_inline(paragraph, text[4:])
            elif marker in ("-", "*") and text.startswith("[x] "):
                paragraph = doc.add_paragraph(style="List Bullet" if indent < 2 else "List Bullet 2")
                run = paragraph.add_run("☑ ")
                set_run_font(run, FONT_SYMBOL, 10.5, color=BRAND)
                add_inline(paragraph, text[4:])
            elif marker in ("-", "*"):
                style = "List Bullet" if indent < 2 else ("List Bullet 2" if indent < 4 else "List Bullet 3")
                paragraph = doc.add_paragraph(style=style)
                add_inline(paragraph, text)
            else:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.75 + indent * 0.2)
                paragraph.paragraph_format.first_line_indent = Cm(-0.75)
                paragraph.paragraph_format.space_after = Pt(2)
                run = paragraph.add_run(f"{marker} ")
                set_run_font(run, FONT_BODY, 10.5, color=TEXT)
                add_inline(paragraph, text)
            index += 1
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()
    return doc


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for source, output_name, title, subtitle in DOCS:
        source_path = ROOT / source
        if not source_path.exists():
            print(f"missing: {source}")
            continue
        markdown = source_path.read_text(encoding="utf-8")
        document = build(markdown, title, subtitle)
        output_path = OUT_DIR / output_name
        try:
            document.save(output_path)
            print(f"built: {output_path.relative_to(ROOT)}")
        except PermissionError:
            print(f"locked, skipped: {output_path.name}（关掉 Word 后重跑）")


if __name__ == "__main__":
    main()
